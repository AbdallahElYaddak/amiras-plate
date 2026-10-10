"""Build the Google Sheet starter file for Amira from the site's /data/*.csv files.

Run from the project folder:  python tools/build_sheet_template.py
Writes sheet-template/AmirasPlate-Menu.xlsx (import it into Google Sheets: File > Import).
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'site', 'data')
OUT = os.path.join(ROOT, 'sheet-template', 'AmirasPlate-Menu.xlsx')

TABS = [  # (csv name, tab title, column widths)
    ('items', 'الأصناف', [16, 40, 16, 10, 8]),
    ('offers', 'العروض', [26, 10, 34, 44, 30, 8]),
    ('categories', 'الأقسام', [16, 46, 20, 10]),
]
HEAD_FILL = PatternFill('solid', fgColor='D3123E')
HEAD_FONT = Font(bold=True, color='FFFFFF')

HELP = [
    'طريقة تعديل المنيو',
    '',
    'الأصناف: كل سطر صنف. اكتبي اسم القسم زي ما هو مكتوب في تبويب "الأقسام" بالظبط.',
    'السعر: رقم بس (من غير "ج"). لو السعر بيتغير، سيبي الخانة فاضية وهيظهر "حسب سعر اليوم".',
    'متاح: "نعم" أو "لا". لو "لا" الصنف هيظهر رمادي ومكتوب جنبه "مش متاح دلوقتي".',
    'العروض: كل اختيار في سطر لوحده جوه نفس الخانة (Ctrl+Enter في جوجل شيت، أو Alt+Enter على الكمبيوتر).',
    'العروض اللي "متاح" فيها "لا" بتختفي من الموقع.',
    'الأقسام: ترتيب السطور هنا هو ترتيب الأقسام في الموقع. الصورة لازم تكون اسم ملف موجود في الموقع.',
    'ممنوع تغيير أسماء الأعمدة في السطر الأول.',
    'التغييرات بتظهر على الموقع خلال حوالي 5 دقايق.',
]


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for name, title, widths in TABS:
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        with open(os.path.join(DATA, name + '.csv'), encoding='utf-8-sig', newline='') as f:
            rows = list(csv.reader(f))
        for r in rows:
            ws.append([int(v) if v.isdigit() else v for v in r])
        for c in ws[1]:
            c.fill, c.font = HEAD_FILL, HEAD_FONT
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = 'A2'
        head = rows[0]
        if 'متاح' in head:
            col = chr(65 + head.index('متاح'))
            dv = DataValidation(type='list', formula1='"نعم,لا"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f'{col}2:{col}500')
    ws = wb.create_sheet('طريقة الاستخدام')
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions['A'].width = 110
    for line in HELP:
        ws.append([line])
    ws['A1'].font = Font(bold=True, size=14, color='D3123E')
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    tmp = OUT + '.tmp.xlsx'
    wb.save(tmp)
    assert os.path.getsize(tmp) > 0
    os.replace(tmp, OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
